"""Input validation and bounded PDF extraction, without Streamlit side effects."""
from io import BytesIO
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from zipfile import ZipFile
from docx import Document

MAX_PDF_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_PAGE_CONTENT_BYTES = 2 * 1024 * 1024
MAX_EXTRACTED_CHARS = 100_000
MIN_TEXT_CHARS = 100
MAX_RESUME_CHARS = 24_000
MAX_JOB_CHARS = 16_000


def extract_docx_text(file_bytes: bytes) -> str:
    """Read paragraphs and tables from a bounded DOCX file."""
    if len(file_bytes) > MAX_PDF_BYTES:
        raise ValueError("The DOCX is over 5 MB. Use a smaller file or paste its text.")
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(m.file_size for m in members) > 25_000_000:
                raise ValueError("The DOCX is too complex. Paste its text instead.")
            if 'word/document.xml' not in archive.namelist():
                raise ValueError("This is not a DOCX document. Paste its text instead.")
        document = Document(BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
        text = '\n'.join(parts).strip()
        if len(text) > MAX_EXTRACTED_CHARS:
            raise ValueError("The DOCX contains too much text. Paste the relevant text.")
        if not text:
            raise ValueError("No readable text found in the DOCX. Paste its text instead.")
        return text
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The DOCX could not be read. Try another file or paste its text.") from exc

def extract_pdf_text(file_bytes: bytes) -> str:
    """Return selectable PDF text, or raise ValueError with a useful message."""
    if len(file_bytes) > MAX_PDF_BYTES:
        raise ValueError("The PDF is over 5 MB. Use a smaller PDF or paste its text.")
    if not file_bytes.startswith(b"%PDF-"):
        raise ValueError("This does not appear to be a PDF. Upload a PDF or paste the text.")

    try:
        reader = PdfReader(BytesIO(file_bytes), strict=False)
        if reader.is_encrypted:
            raise ValueError("This PDF is password-protected. Unlock it or paste the text.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("The PDF has over 20 pages. Upload a shorter resume or paste its text.")
        pages = []
        total_chars = 0
        for page in reader.pages:
            content = page.get_contents()
            if content is not None and len(content.get_data()) > MAX_PAGE_CONTENT_BYTES:
                raise ValueError("A PDF page is too complex. Export a simpler PDF or paste its text.")
            page_text = page.extract_text() or ""
            total_chars += len(page_text)
            if total_chars > MAX_EXTRACTED_CHARS:
                raise ValueError("The PDF contains too much text. Upload a shorter resume or paste its text.")
            pages.append(page_text)
    except ValueError:
        raise
    except (PdfReadError, OSError, EOFError, TypeError, KeyError, IndexError) as exc:
        raise ValueError("The PDF could not be read. Try another PDF or paste the text.") from exc
    except Exception as exc:
        # PDF parsing can fail in several ways, depending on the malformed file.
        raise ValueError("The PDF could not be read. Try another PDF or paste the text.") from exc

    text = "\n\n".join(page.strip() for page in pages if page.strip()).strip()
    if not text:
        raise ValueError(
            "No selectable text was found. This may be a scanned image; paste the text instead."
        )
    return text


def validate_text(resume: str, job: str) -> list[str]:
    """Return input problems; never call an LLM when this list is nonempty."""
    problems = []
    if len(resume.strip()) < MIN_TEXT_CHARS:
        problems.append("Resume: enter at least 100 characters of readable text.")
    if len(job.strip()) < MIN_TEXT_CHARS:
        problems.append("Job description: enter at least 100 characters.")
    return problems



def trim_text(text: str, limit: int) -> str:
    text = text.strip()
    if len(text) <= limit:
        return text
    head = text[:limit]
    boundary = max(head.rfind("\n\n"), head.rfind("\n"), head.rfind(" "))
    if boundary >= int(limit * 0.8):
        head = head[:boundary]
    return head.strip()
