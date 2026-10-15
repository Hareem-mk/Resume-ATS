"""Offline checks; no real Groq calls or API key needed."""
import json
import types
import unittest
from io import BytesIO
from unittest.mock import Mock, patch
from docx import Document
from inputs import extract_docx_text
from review import ReviewError, build_quick_crew, parse_output, render_report, review_quick

RESUME = 'Participated in HAZOP workshops. Reduced inspection time by 10%. ' * 3
JOB = 'Must have HAZOP experience. Nice to have SIL certification. ' * 3


def output_data():
    return {'requirements': [
        {'requirement': 'HAZOP', 'priority': 'Must-have', 'status': 'Matched',
         'job_evidence': 'Must have HAZOP experience.', 'evidence': 'Participated in HAZOP workshops.'},
        {'requirement': 'SIL', 'priority': 'Nice-to-have', 'status': 'Missing',
         'job_evidence': 'Nice to have SIL certification.', 'evidence': 'Not found in resume'}],
        'recommendations': [{'priority': 'High', 'action': 'Clarify HAZOP responsibilities.'}],
        'rewrites': [], 'ats_keywords': ['SIL', 'invented-keyword']}


class ReviewTests(unittest.TestCase):
    def test_grounded_score_and_headings(self):
        result = render_report(output_data(), RESUME, JOB)
        self.assertEqual(result['score'], 67)
        for n in range(1, 9):
            self.assertIn(f'## {n}.', result['markdown'])
        self.assertNotIn('invented-keyword', result['markdown'])
        self.assertIn('| --- | --- | --- |\n|', result['markdown'])

    def test_fabricated_resume_quote_is_downgraded(self):
        data = output_data()
        data['requirements'][0]['evidence'] = 'I invented this experience'
        self.assertEqual(render_report(data, RESUME, JOB)['score'], 0)

    def test_unverified_job_requirements_fail(self):
        data = output_data()
        for row in data['requirements']:
            row['job_evidence'] = 'not in this job'
        with self.assertRaises(ReviewError):
            render_report(data, RESUME, JOB)

    def test_rewrite_numbers_must_come_from_before(self):
        data = output_data()
        data['rewrites'] = [{'before': 'Participated in HAZOP workshops.',
                             'after': 'Led 10 HAZOP workshops.'}]
        result = render_report(data, RESUME, JOB)
        self.assertNotIn('Led 10', result['markdown'])

    def test_invalid_json(self):
        with self.assertRaises(ReviewError):
            parse_output('not JSON')
        self.assertEqual(parse_output('```json\n{"ok": true}\n```'), {'ok': True})

    def test_short_input_never_calls_provider(self):
        with patch('review.build_quick_crew') as mocked:
            with self.assertRaises(ReviewError):
                review_quick('short', JOB, 'test-key', 'test-model')
            mocked.assert_not_called()

    def test_rate_limit_has_three_total_attempts(self):
        error = Exception('rate limit'); error.status_code = 429
        with patch('review.build_quick_crew', side_effect=error) as mocked, patch('review.time.sleep') as sleep:
            with self.assertRaisesRegex(ReviewError, 'rate limit'):
                review_quick(RESUME, JOB, 'test-key', 'test-model')
            self.assertEqual(mocked.call_count, 3)
            self.assertEqual([c.args[0] for c in sleep.call_args_list], [1, 2])

    def test_bad_key_does_not_retry(self):
        error = Exception('private-provider-payload'); error.status_code = 401
        with patch('review.build_quick_crew', side_effect=error) as mocked:
            with self.assertRaises(ReviewError) as raised:
                review_quick(RESUME, JOB, 'test-key', 'test-model')
            self.assertEqual(mocked.call_count, 1)
            self.assertNotIn('private-provider-payload', str(raised.exception))

    def test_quick_crew_has_one_agent_and_task(self):
        fake = types.ModuleType('crewai')
        fake.LLM, fake.Agent, fake.Task, fake.Crew = Mock(), Mock(), Mock(), Mock()
        fake.Process = types.SimpleNamespace(sequential='sequential')
        with patch.dict('sys.modules', {'crewai': fake}):
            build_quick_crew(RESUME, JOB, 'test-key', 'openai/gpt-oss-120b')
        self.assertEqual(fake.LLM.call_args.kwargs['model'], 'groq/openai/gpt-oss-120b')
        self.assertEqual(fake.LLM.call_args.kwargs['temperature'], 0.2)
        self.assertEqual(len(fake.Crew.call_args.kwargs['agents']), 1)
        self.assertEqual(len(fake.Crew.call_args.kwargs['tasks']), 1)
        self.assertFalse(fake.Crew.call_args.kwargs['memory'])

    def test_docx_paragraphs_and_tables(self):
        doc = Document(); doc.add_paragraph('Resume experience')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'HAZOP'; table.cell(0, 1).text = 'SIMOPS'
        buffer = BytesIO(); doc.save(buffer)
        text = extract_docx_text(buffer.getvalue())
        self.assertIn('Resume experience', text)
        self.assertIn('HAZOP | SIMOPS', text)
        with self.assertRaises(ValueError):
            extract_docx_text(b'broken DOCX')


if __name__ == '__main__':
    unittest.main()
